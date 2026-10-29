from io import BytesIO

from docx import Document as DocxDocument
from pptx import Presentation

from extractors.clean import clean_native


def extract_pptx(file_bytes):
    try:
        prs = Presentation(BytesIO(file_bytes))
        partes = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    partes.append(shape.text)
            if slide.has_notes_slide:
                notas = slide.notes_slide.notes_text_frame.text
                if notas:
                    partes.append(notas)
        return clean_native("\n".join(partes))
    except Exception:
        return ""


def extract_docx(file_bytes):
    try:
        doc = DocxDocument(BytesIO(file_bytes))
        partes = [p.text for p in doc.paragraphs]
        return clean_native("\n".join(partes))
    except Exception:
        return ""

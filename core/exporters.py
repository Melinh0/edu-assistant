import os
import subprocess
import tempfile
from io import BytesIO

from docx import Document
from docx.shared import Pt
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

MIME = {
    "txt": "text/plain",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


def pandoc_convert(text, output_format, extra_args=None):
    extra_args = extra_args or []
    for engine in ("xelatex", "pdflatex"):
        md_path = None
        out_path = None
        try:
            with tempfile.NamedTemporaryFile(mode="w", suffix=".md", delete=False, encoding="utf-8") as md_file:
                md_file.write(text)
                md_path = md_file.name
            suffix = ".docx" if output_format == "docx" else ".pdf"
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as out_file:
                out_path = out_file.name
            cmd = ["pandoc", md_path, "-o", out_path, "--standalone"]
            if output_format == "pdf":
                cmd.extend(["--pdf-engine", engine])
            cmd.extend(extra_args)
            subprocess.run(cmd, check=True, capture_output=True, text=True)
            with open(out_path, "rb") as f:
                return f.read()
        except (subprocess.CalledProcessError, OSError, FileNotFoundError):
            continue
        finally:
            for path in (md_path, out_path):
                if path and os.path.exists(path):
                    os.unlink(path)
    return None


def _docx_fallback(text):
    doc = Document()
    for line in text.split("\n"):
        if line.strip():
            p = doc.add_paragraph(line)
            p.style.font.size = Pt(12)
        else:
            doc.add_paragraph()
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _pdf_fallback(text):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    style_normal = styles["Normal"]
    style_normal.fontName = "Helvetica"
    style_normal.fontSize = 12
    style_normal.leading = 14
    style_normal.alignment = TA_LEFT
    story = []
    for line in text.split("\n"):
        if line.strip() == "":
            story.append(Spacer(1, 12))
        else:
            story.append(Paragraph(line, style_normal))
    doc.build(story)
    buffer.seek(0)
    return buffer


def create_download_link(text, filename, formato):
    if formato == "txt":
        buf = BytesIO(text.encode("utf-8"))
        return buf, f"{filename}.txt", MIME["txt"]

    if formato == "docx":
        data = pandoc_convert(text, "docx")
        if data:
            buf = BytesIO(data)
            buf.seek(0)
            return buf, f"{filename}.docx", MIME["docx"]
        return _docx_fallback(text), f"{filename}.docx", MIME["docx"]

    if formato == "pdf":
        data = pandoc_convert(text, "pdf")
        if data:
            buf = BytesIO(data)
            buf.seek(0)
            return buf, f"{filename}.pdf", MIME["pdf"]
        return _pdf_fallback(text), f"{filename}.pdf", MIME["pdf"]

    raise ValueError(f"Formato não suportado: {formato}")


def save_uploaded_file(uploaded_file, save_dir):
    os.makedirs(save_dir, exist_ok=True)
    file_path = os.path.join(save_dir, uploaded_file.name)
    with open(file_path, "wb") as f:
        f.write(uploaded_file.getbuffer())
    return file_path
